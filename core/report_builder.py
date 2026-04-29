"""
Geração do relatório Excel com 6 abas.
"""
from __future__ import annotations
import datetime
import io
from decimal import Decimal
from typing import Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from .normalize import (
    STATUS_CONCILIADO, STATUS_CONCILIADO_MANUAL,
    STATUS_REVISAR, STATUS_REVISAR_COLISAO,
    STATUS_IGNORADO_SEM_PAR, STATUS_SEM_PAREAMENTO,
    STATUS_IGNORADO_USUARIO,
)

# Cores
COR_CONCILIADO     = "C6EFCE"  # verde claro
COR_MANUAL         = "CCFFCC"  # verde mais claro
COR_SEM_PAR        = "FFEB9C"  # amarelo
COR_REVISAR        = "FFC7CE"  # vermelho claro
COR_IGNORADO       = "D9D9D9"  # cinza
COR_CABECALHO      = "2F75B6"  # azul

_DATE_FMT = "DD/MM/YYYY"


def _fill(hex_color: str) -> PatternFill:
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def _header_font() -> Font:
    return Font(bold=True, color="FFFFFF")


def _auto_width(ws, min_width=10, max_width=60):
    for col in ws.columns:
        length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(length + 2, min_width), max_width)


def _apply_date_format(ws) -> None:
    """
    Aplica number_format DD/MM/YYYY a cada célula de data na última linha gravada.
    Gravar datetime.date diretamente (sem strftime) garante tipo Date no Excel,
    o que habilita filtro hierárquico Ano > Mês > Dia e ordena corretamente.
    """
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, (datetime.date, datetime.datetime)):
            cell.number_format = _DATE_FMT


def _as_date(value) -> object:
    """Retorna datetime.date se o valor for um objeto de data; caso contrário devolve o valor original."""
    if isinstance(value, datetime.datetime):
        return value.date()
    return value


def _to_float(v) -> object:
    """Converte Decimal para float; passa outros valores sem alteração."""
    return float(v) if isinstance(v, Decimal) else v


def build_report(
    df_bnk: pd.DataFrame,
    df_fin: pd.DataFrame,
    depara: Optional[dict] = None,
) -> bytes:
    """
    Constrói workbook Excel com 6 abas e retorna bytes.
    depara: dict {classif -> (debito, credito)}
    """
    for col in ["_id", "_data", "_valor", "_historico", "_classif", "_status", "_metodo", "_id_bnk"]:
        if col not in df_fin.columns:
            df_fin[col] = ""

    for col in ["_id", "_data", "_valor", "_historico", "_classif", "_status", "_metodo", "_ids_fin"]:
        if col not in df_bnk.columns:
            df_bnk[col] = ""

    depara = depara or {}
    wb = Workbook()
    wb.remove(wb.active)

    _build_consolidado(wb, df_bnk, df_fin, depara)
    _build_extrato(wb, df_bnk)
    _build_financeiro(wb, df_fin)
    _build_sem_par_bnk(wb, df_bnk)
    _build_sem_par_fin(wb, df_fin)
    _build_resumo(wb, df_bnk, df_fin)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _resolve_classif(ids_fin_str: str, fin_by_id: dict) -> str:
    if not ids_fin_str:
        return ""
    ids = [x.strip() for x in str(ids_fin_str).split(";") if x.strip()]
    classifs = []
    for id_f in ids:
        rec = fin_by_id.get(id_f)
        if rec:
            c = str(rec.get("_classif", "")).strip()
            if c and c not in classifs:
                classifs.append(c)
    if len(classifs) == 0:
        return ""
    if len(classifs) == 1:
        return classifs[0]
    return "[MULTIPLAS: " + ", ".join(classifs) + "]"


def _resolve_historico_fin(ids_fin_str: str, fin_by_id: dict, sep: str = " - ") -> str:
    if not ids_fin_str:
        return ""
    ids = [x.strip() for x in str(ids_fin_str).split(";") if x.strip()]
    hists = []
    for id_f in ids:
        rec = fin_by_id.get(id_f)
        if rec:
            h = str(rec.get("_historico", "")).strip()
            if h and h not in hists:
                hists.append(h)
    return sep.join(hists)


def _build_consolidado(wb, df_bnk, df_fin, depara):
    ws = wb.create_sheet("Relatorio Consolidado")
    headers = [
        "Data", "Historico", "Valor", "Classificacao Financeira",
        "Historico Financeiro",
        "Tipo Conciliacao", "Debito (Conta)", "Credito (Conta)",
        "ID Banco", "ID Financeiro", "Metodo", "Status",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _fill(COR_CABECALHO)
        cell.font = _header_font()
        cell.alignment = Alignment(horizontal="center")

    status_cor = {
        STATUS_CONCILIADO:        COR_CONCILIADO,
        STATUS_CONCILIADO_MANUAL: COR_MANUAL,
        STATUS_SEM_PAREAMENTO:    COR_SEM_PAR,
        STATUS_REVISAR:           COR_REVISAR,
        STATUS_REVISAR_COLISAO:   COR_REVISAR,
        STATUS_IGNORADO_SEM_PAR:  COR_IGNORADO,
        STATUS_IGNORADO_USUARIO:  COR_IGNORADO,
    }

    # Pre-indexa financeiros por _id — elimina scans lineares O(n) por linha
    fin_by_id: dict = {
        rec["_id"]: rec
        for rec in df_fin[["_id", "_valor", "_historico", "_classif"]].to_dict("records")
    }

    for row in df_bnk.to_dict("records"):
        metodo = str(row.get("_metodo", ""))
        ids_fin_str = str(row.get("_ids_fin", ""))
        status = str(row.get("_status", ""))
        cor = status_cor.get(status, "FFFFFF")

        data_val = _as_date(row.get("_data", ""))

        is_expandable = (
            ("1:N" in metodo and "ambiguo" not in metodo)
            or metodo == "manual"
        ) and bool(ids_fin_str.strip())

        if is_expandable:
            tipo_expand = "Desmembrado (1:N)" if "1:N" in metodo else "Manual"
            ids_fin = [x.strip() for x in ids_fin_str.split(";") if x.strip()]
            for id_f in ids_fin:
                fin_row = fin_by_id.get(id_f)
                if fin_row is None:
                    continue

                classif = str(fin_row.get("_classif", "")).strip()
                debito = credito = ""
                if classif and not classif.startswith("[MULTIPLAS"):
                    par = depara.get(classif, ("", ""))
                    debito, credito = par[0], par[1]

                valor_val = _to_float(fin_row.get("_valor", ""))
                hist_fin = str(fin_row.get("_historico", "")).strip()

                linha = [
                    data_val,
                    str(row.get("_historico", "")),
                    valor_val,
                    classif,
                    hist_fin,
                    tipo_expand,
                    debito,
                    credito,
                    str(row.get("_id", "")),
                    id_f,
                    metodo,
                    status,
                ]
                ws.append(linha)
                _apply_date_format(ws)
                for cell in ws[ws.max_row]:
                    cell.fill = _fill(cor)
        else:
            classif = _resolve_classif(ids_fin_str, fin_by_id)
            debito = credito = ""
            if classif and not classif.startswith("[MULTIPLAS"):
                par = depara.get(classif, ("", ""))
                debito, credito = par[0], par[1]

            tipo = ""
            if "1:1" in metodo:
                tipo = "Exato"
            elif "N:1" in metodo:
                tipo = "Agrupado (N:1)"
            elif "1:N" in metodo:
                tipo = "Desmembrado (1:N)"
            elif "manual" in metodo:
                tipo = "Manual"

            hist_fin = _resolve_historico_fin(ids_fin_str, fin_by_id)
            valor_val = _to_float(row.get("_valor", ""))

            linha = [
                data_val,
                str(row.get("_historico", "")),
                valor_val,
                classif,
                hist_fin,
                tipo,
                debito,
                credito,
                str(row.get("_id", "")),
                ids_fin_str,
                metodo,
                status,
            ]
            ws.append(linha)
            _apply_date_format(ws)
            for cell in ws[ws.max_row]:
                cell.fill = _fill(cor)

    _auto_width(ws)


def _build_extrato(wb, df_bnk):
    ws = wb.create_sheet("Extrato Bancario")
    cols_base = ["_id", "_data", "_valor", "_historico", "_status", "_metodo", "_ids_fin"]
    extras = [c for c in df_bnk.columns if not c.startswith("_")]
    headers = cols_base + extras
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _fill(COR_CABECALHO)
        cell.font = _header_font()
    for row in df_bnk.to_dict("records"):
        linha = []
        for c in headers:
            v = row.get(c, "")
            if isinstance(v, Decimal):
                v = float(v)
            elif isinstance(v, datetime.datetime):
                v = v.date()
            linha.append(v)
        ws.append(linha)
        _apply_date_format(ws)
    _auto_width(ws)


def _build_financeiro(wb, df_fin):
    ws = wb.create_sheet("Financeiro")
    cols_base = ["_id", "_data", "_valor", "_historico", "_classif", "_status", "_metodo", "_id_bnk"]
    extras = [c for c in df_fin.columns if not c.startswith("_")]
    headers = cols_base + extras
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _fill(COR_CABECALHO)
        cell.font = _header_font()
    for row in df_fin.to_dict("records"):
        linha = []
        for c in headers:
            v = row.get(c, "")
            if isinstance(v, Decimal):
                v = float(v)
            elif isinstance(v, datetime.datetime):
                v = v.date()
            linha.append(v)
        ws.append(linha)
        _apply_date_format(ws)
    _auto_width(ws)


def _build_sem_par_bnk(wb, df_bnk):
    ws = wb.create_sheet("Banco Sem Par")
    headers = ["ID Banco", "Data", "Valor", "Historico", "Status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _fill(COR_CABECALHO)
        cell.font = _header_font()
    revisar_set = {STATUS_SEM_PAREAMENTO, STATUS_REVISAR, STATUS_REVISAR_COLISAO}
    for row in df_bnk.to_dict("records"):
        if str(row.get("_status", "")) not in revisar_set:
            continue
        data_val = _as_date(row.get("_data", ""))
        valor_val = _to_float(row.get("_valor", ""))
        ws.append([
            str(row.get("_id", "")),
            data_val,
            valor_val,
            str(row.get("_historico", "")),
            str(row.get("_status", "")),
        ])
        _apply_date_format(ws)
        for cell in ws[ws.max_row]:
            cell.fill = _fill(COR_SEM_PAR)
    _auto_width(ws)


def _build_sem_par_fin(wb, df_fin):
    ws = wb.create_sheet("Financeiro Sem Par")
    headers = ["ID Fin", "Data", "Valor", "Historico", "Classificacao", "Status"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _fill(COR_CABECALHO)
        cell.font = _header_font()
    revisar_set = {STATUS_IGNORADO_SEM_PAR, STATUS_REVISAR, STATUS_REVISAR_COLISAO}
    for row in df_fin.to_dict("records"):
        if str(row.get("_status", "")) not in revisar_set:
            continue
        data_val = _as_date(row.get("_data", ""))
        valor_val = _to_float(row.get("_valor", ""))
        ws.append([
            str(row.get("_id", "")),
            data_val,
            valor_val,
            str(row.get("_historico", "")),
            str(row.get("_classif", "")),
            str(row.get("_status", "")),
        ])
        _apply_date_format(ws)
        for cell in ws[ws.max_row]:
            cell.fill = _fill(COR_SEM_PAR)
    _auto_width(ws)


def _build_resumo(wb, df_bnk, df_fin):
    ws = wb.create_sheet("Resumo")
    ws.append(["Metrica", "Quantidade", "Valor Total"])
    for cell in ws[1]:
        cell.fill = _fill(COR_CABECALHO)
        cell.font = _header_font()

    # Converte coluna _valor para float uma única vez (evita loop com Decimal por linha).
    # .astype(float) garante dtype numérico mesmo em DataFrames vazios (apply preserva dtype str).
    bnk_vals = df_bnk["_valor"].apply(lambda v: float(v) if v else 0.0).astype(float)
    fin_vals = df_fin["_valor"].apply(lambda v: float(v) if v else 0.0).astype(float)

    def soma(status_series, vals_series, status_list):
        mask = status_series.isin(status_list)
        return int(mask.sum()), round(float(vals_series[mask].sum()), 2)

    metricas = [
        ("Conciliados (auto)",   df_bnk["_status"], bnk_vals, [STATUS_CONCILIADO]),
        ("Conciliados (manual)", df_bnk["_status"], bnk_vals, [STATUS_CONCILIADO_MANUAL]),
        ("A Revisar",            df_bnk["_status"], bnk_vals, [STATUS_REVISAR, STATUS_REVISAR_COLISAO]),
        ("Banco Sem Par",        df_bnk["_status"], bnk_vals, [STATUS_SEM_PAREAMENTO]),
        ("Financeiro Sem Par",   df_fin["_status"], fin_vals, [STATUS_IGNORADO_SEM_PAR]),
        ("Ignorados (usuario)",  df_bnk["_status"], bnk_vals, [STATUS_IGNORADO_USUARIO]),
    ]

    for label, status_series, vals_series, statuses in metricas:
        qtd, total = soma(status_series, vals_series, statuses)
        ws.append([label, qtd, total])

    _auto_width(ws)
